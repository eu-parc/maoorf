#!/usr/bin/env python3
"""Convert a filled-in OECD Omics Reporting Framework (OORF) Excel workbook into a LinkML-compliant JSON/YAML instance."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl
import yaml


def clean_val(val: Any) -> str | None:
    """Convert Excel cell value to a stripped string, or None if empty."""
    if val is None:
        return None
    s = str(val).strip()
    return s if s != "" else None


def get_row_input(ws: openpyxl.worksheet.worksheet.Worksheet, row_idx: int, col_idx: int = 4) -> str | None:
    """Read the INPUT column (default Column D = 4) from a worksheet row."""
    return clean_val(ws.cell(row=row_idx, column=col_idx).value)


def prune_empty(obj: Any) -> Any:
    """Recursively remove keys with None values or empty dicts."""
    if isinstance(obj, dict):
        cleaned = {}
        for k, v in obj.items():
            pv = prune_empty(v)
            if pv is not None and pv != {} and pv != []:
                cleaned[k] = pv
        return cleaned
    if isinstance(obj, list):
        return [prune_empty(x) for x in obj if prune_empty(x) not in (None, {}, [])]
    return obj


def parse_ssrm(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Parse the SSRM sheet (1-indexed row numbers, same in the EUT046 and RH010 example templates)."""
    r = lambda row: get_row_input(ws, row, 4)
    return prune_empty({
        "study_identifiers": {
            "abstract": r(6),
            "study_identifier": r(8),
            "standardised_toxicology_dataset": r(9),
            "complete_omics_dataset": r(10),
        },
        "study_rationale": {
            "background_information": r(12),
            "objectives": r(13),
            "test_item": {
                "test_item_name": r(15),
                "test_item_cas": r(16),
                "test_item_smiles": r(17),
                "test_item_iupac_name": r(18),
                "test_item_additional_information": r(19),
                "test_item_sources_of_identifiers": r(20),
            },
            "test_system_characteristics_in_vivo": {
                "species": r(22),
                "strain": r(23),
                "sex": r(24),
                "age": r(25),
            },
            "test_system_characteristics_in_vitro": {
                "cell_type": r(27),
                "species_of_origin": r(28),
            },
            "study_design": {
                "dose_levels": r(30),
                "dose_intervals": r(31),
                "number_of_biological_replicates_per_treatment_condition": r(32),
                "timetable_sample_collection": r(33),
            },
            "treatment_conditions": {
                "route_of_administration": r(35),
                "identity_of_vehicle": r(36),
            },
            "sample_collection": {
                "type_of_biological_sample_for_omics": r(38),
                "sample_preservation_method": r(39),
            },
        },
        "data_acquisition_and_processing": {
            "daprm_modules": r(41),
            "omics_technology_manufacturer_and_model": r(42),
            "description_and_rationale_for_data_acquisition_approach": r(43),
            "sample_preparation_method": r(44),
            "demonstration_of_quality_of_analysis": r(45),
        },
        "data_analysis": {
            "darm_modules": r(47),
            "description_and_rationale_for_data_analysis_approach": r(48),
        },
    })


def parse_term(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Parse the TERM sheet (1-indexed row numbers, same in the EUT046 and RH010 example templates)."""
    r = lambda row: get_row_input(ws, row, 4)
    return prune_empty({
        "study_rationale": {
            "background_information": r(6),
            "objectives": r(7),
            "test_guideline_compliance": r(8),
            "mechanistic_understanding": r(9),
            "model_selection": {
                "in_vivo_species_model_and_tissue_rationale": r(11),
                "species_and_strain_rationale": r(12),
                "in_vitro_relevance": r(13),
            },
            "dose_concentration_level_and_interval_selection": r(14),
            "route_of_administration": r(15),
            "time_point_selection": r(16),
            "samples_and_replicates": {
                "biological_replicate_number": r(18),
                "number_of_technical_and_analytical_replicates": r(19),
            },
            "limitations": r(20),
        },
        "test_and_control_items": {
            "test_item_name": r(21),
            "formulation_composition": {
                "substances_in_mixture": r(23),
                "relative_proportions_of_substances": r(24),
            },
            "preparation_of_test_item_composition": {
                "concentration_of_test_item": r(26),
                "concentration_of_diluents": r(27),
                "identification_of_impurities": r(28),
            },
            "physicochemical_properties": {
                "appearance_physical_state_colour": r(30),
                "molecular_weight": r(31),
                "melting_point_freezing_point": r(32),
                "boiling_point": r(33),
                "ph": r(34),
                "viscosity": r(35),
                "density": r(36),
                "vapour_pressure": r(37),
                "partition_coefficient": r(38),
                "water_solubility": r(39),
                "fat_solubility": r(40),
                "particle_size_distribution": r(41),
                "additional_physicochemical_information": r(42),
            },
            "chemical_stability": {
                "stability_in_organic_solvents": r(44),
                "stability_thermal_sunlight_metals": r(45),
                "stability_dissociation_constant": r(46),
            },
            "commercial_source": {
                "vendor": r(48),
                "manufacture_id": r(49),
                "lot_batch_number": r(50),
                "purity": r(51),
                "salt_form": r(52),
                "expiration_date": r(53),
                "storage": r(54),
            },
            "test_item_specific_identifiers": {
                "cas": r(56),
                "smiles": r(57),
                "iupac_name": r(58),
                "additional_information": r(59),
                "sources_of_identifiers": r(60),
            },
        },
        "test_system_characteristics": {
            "general_characteristics_in_vivo": {
                "animal_species": r(62),
                "strain": r(63),
                "sex": r(64),
                "age_during_study": r(65),
                "developmental_stage": r(66),
                "individual_weights_lengths_at_start": r(67),
                "supplier": r(68),
                "interventions_before_or_during_experiment": r(69),
                "quality_criteria_before_use": r(70),
                "health_status_and_acclimation": r(71),
                "randomisation_of_animals_to_groups": r(72),
                "other": r(73),
            },
            "general_characteristics_in_vitro": {
                "cell_type": r(75),
                "origin": r(76),
                "cell_passage_number": r(77),
                "differentiation_stage": r(78),
                "absence_of_mycoplasma": r(79),
                "metabolic_competence": r(80),
                "supplier": r(81),
                "quality_criteria_before_use": r(82),
                "other": r(83),
            },
            "housing_and_husbandry_in_vivo": {
                "type_of_facility": r(85),
                "type_of_cage_or_housing": r(86),
                "bedding_material": r(87),
                "number_of_cage_companions": r(88),
                "tank_shape_and_material": r(89),
                "breeding_programme": r(90),
                "light_dark_cycle": r(91),
                "temperature": r(92),
                "quality_of_water": r(93),
                "type_of_food": r(94),
                "access_to_food_and_water": r(95),
                "environmental_enrichment": r(96),
                "methods_for_fertilisation_collection_of_eggs": r(97),
                "other": r(98),
            },
            "cell_culture_conditions_in_vitro": {
                "incubation_characteristics": {
                    "temperature": r(101),
                    "co2_o2_conditions": r(102),
                    "humidity": r(103),
                    "other": r(104),
                },
                "state_of_cells_before_use": {
                    "viability": r(106),
                    "quality_control": r(107),
                    "morphology": r(108),
                    "recommended_confluency_of_use": r(109),
                    "other": r(110),
                },
                "culture_media": {
                    "use_of_serum": r(112),
                    "use_of_antibiotics": r(113),
                },
                "growth_support_substrates": {
                    "use_of_feeder_cells": r(115),
                    "use_of_matrixes_or_scaffolds": r(116),
                },
            },
        },
        "study_design": {
            "three_rs_considerations": {
                "three_rs_description": r(118),
            },
            "dose_selection": {
                "dose_levels": r(120),
                "dose_intervals": r(121),
            },
            "test_method_instruments_equipment_and_reagents": {
                "commercial_source": r(123),
                "manufacturers_instrument_model_identification": r(124),
                "manufacturers_reagent_and_kit_information": r(125),
                "special_safety_handling_requirements": r(126),
            },
            "types_of_treatments": {
                "controls": r(128),
                "pre_treatments": r(129),
                "acclimation": r(130),
                "types_of_replicates": r(131),
            },
            "number_of_animals_samples_per_treatment": {
                "biological_replicates": r(133),
                "technical_analytical_replicates": r(134),
            },
            "statistical_design": {
                "exposure_design": r(136),
                "sampling_schemes": r(137),
                "sample_blinding": r(138),
            },
            "observations_examinations_during_treatment": {
                "in_life_clinical_observations": r(140),
                "toxicokinetics": r(141),
                "histopathology_and_organ_weight": r(142),
                "clinical_chemistry_and_pathology_in_vivo": r(143),
                "reason_animals_removed": r(144),
                "cytological_analyses_in_vivo": r(145),
                "cytobiological_examinations_in_vitro": r(146),
                "other_biochemical_molecular_analyses": r(147),
            },
        },
        "treatment_conditions": {
            "route_of_administration_in_vivo": r(148),
            "route_of_administration_in_vitro": r(149),
            "housing_condition_modifications_in_vivo": r(150),
            "culture_condition_modifications_in_vitro": r(151),
            "test_item_preparation": {
                "dilution_in_vehicle": r(153),
                "preparation_steps": r(154),
                "separation_steps": r(155),
                "extraction_steps": r(156),
                "storage_conditions": r(157),
                "stability_during_storage": r(158),
                "expiration_date": r(159),
                "nominal_or_measured_concentrations": r(160),
                "analytic_controls_on_measured_concentrations": r(161),
                "dosing_solution_homogeneity": r(162),
            },
            "test_item_stability_reactivity": {
                "stability_under_test_conditions": r(164),
                "solubility_and_stability_in_solvent_vehicle": r(165),
                "reactivity_with_solvent_vehicle_or_medium": r(166),
                "impact_of_separation_and_extraction_steps": r(167),
            },
            "test_item_preparation_in_vivo": {
                "formulation_diet_preparation_procedures": r(169),
                "test_atmosphere_generation_and_chamber_description": r(170),
                "achieved_concentration": r(171),
                "stability_and_homogeneity_of_preparation": r(172),
                "test_material_intake": r(173),
            },
            "vehicle_description_and_delivery_volume": {
                "identity_of_vehicle": r(175),
                "delivery_volume": r(176),
                "final_concentration_of_vehicle": r(177),
                "maximum_volume_of_liquid_administered": r(178),
            },
            "exposure_schedule": {
                "frequency_of_test_item_administration": r(180),
                "time_of_day_of_dosing_in_vivo": r(181),
                "time_of_dosing_after_seeding_in_vitro": r(182),
                "recovery_period": r(183),
            },
            "housing_culture_condition_deviations": {
                "temperature": r(185),
                "humidity": r(186),
                "co2_percentage": r(187),
                "ph": r(188),
                "availability_and_quality_of_nutrients": r(189),
                "other": r(190),
            },
        },
        "study_exit_and_sample_collection": {
            "type_of_biological_sample": r(191),
            "study_exit_in_vivo": {
                "anaesthetic_used": r(193),
                "analgesic_used": r(194),
                "method_of_euthanasia": r(195),
                "phenotypic_characteristics": r(196),
                "methods_used_for_collection_of_biological_samples": r(197),
            },
            "study_exit_in_vitro": {
                "collection_of_biological_material": r(199),
                "cell_density_at_time_of_harvesting": r(200),
                "growth_phase_stage": r(201),
                "number_of_culture_passage": r(202),
                "morphology": r(203),
                "methods_used_for_collection_of_biological_samples": r(204),
            },
            "sampling_vial": {
                "type_of_vial_or_tube": r(206),
                "chemicals_within_sample_tube": r(207),
                "chemicals_added_to_preserve_samples": r(208),
            },
            "washing": {
                "washing_solvents": r(210),
                "washing_procedure": r(211),
            },
            "quenching": {
                "quench_solvent": r(213),
                "quenching_procedure": r(214),
            },
            "pooling_of_samples": {
                "sample_pooling_procedures": r(216),
            },
            "sample_storage_and_transport": {
                "post_sample_collection_handling": r(218),
                "storage_temperature_and_duration": r(219),
                "transportation_method": r(220),
                "number_of_freeze_thaw_cycles": r(221),
            },
            "time_table": {
                "treatments": r(223),
                "sample_collections": r(224),
                "time_since_last_dose_administered": r(225),
                "time_to_sample_extractions": r(226),
            },
        },
        "sample_identification_codes": {
            "lims": r(227),
            "method_schema_for_assigning_unique_sample_codes": r(228),
            "metadata": r(229),
        },
        "supporting_data_streams": {
            "supporting_sources_of_data_or_information": r(230),
        },
    })


def parse_daprm_rna_seq(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Parse the DAPRM - RNA-Seq sheet (1-indexed row numbers, same in the EUT046 and RH010 example templates)."""
    r = lambda row: get_row_input(ws, row, 4)
    return prune_empty({
        "type": "RNASeqDAPRM",
        "technology": {
            "type_and_version_of_sequencing_platform": r(6),
            "size_and_type_of_sequencing": r(7),
            "flow_cell_used": r(8),
            "targeting_probe_annotation": r(9),
            "library_type": r(10),
            "purpose": r(11),
            "other": r(12),
        },
        "transcriptomics_experimental_design": {
            "rna_processing": {
                "rna_extraction": {
                    "type_of_extracted_rna": r(15),
                    "extraction_and_purification_techniques": r(16),
                    "procedures_for_mrna_enrichment": r(17),
                    "storage_conditions": r(18),
                },
                "quantification_and_qualification_of_rna": {
                    "tool_for_rna_assessment": r(20),
                    "rna_quality": r(21),
                    "rna_quantity": r(22),
                },
            },
            "library_preparation": {
                "library_preparation_applied": {
                    "manual_or_automated_system": r(25),
                    "fragmentation_strategy": r(26),
                    "probe_manifest": r(27),
                    "number_of_pcr_amplification_cycles": r(28),
                    "targeting_probe_annealing_ligation_and_pcr_steps": r(29),
                    "other": r(30),
                },
            },
            "sample_pooling_protocol": {
                "sample_pooling_occurred": r(32),
                "indexing_strategy": r(33),
                "number_of_samples_pooled_per_sequencing_unit": r(34),
                "table_of_barcodes_assigned_to_each_sample": r(35),
                "other": r(36),
            },
            "batch_processing_of_samples": {
                "number_of_batches": r(38),
                "number_of_samples_in_each_batch": r(39),
                "method_of_sample_assignment_to_batches": r(40),
            },
            "wet_lab_quality_control": {
                "use_of_chip_based_electrophoresis": r(42),
                "use_of_spike_in_controls": {
                    "type_of_spike_in_used": r(44),
                    "amount_of_sequence_used": r(45),
                    "source_of_spike_in": r(46),
                    "step_of_introduction": r(47),
                },
                "use_of_unique_molecular_identifiers": {
                    "type_of_umis": r(49),
                    "source_of_umis": r(50),
                },
                "other": r(51),
            },
            "sequencing_quality_control_criteria": {
                "quality_control_standard": {
                    "type_of_standard": r(54),
                    "source_of_standard": r(55),
                    "quantity_used": r(56),
                },
                "sequencing_quality_metrics": {
                    "number_of_clusters_passing_filtering": r(58),
                    "average_quality_score": r(59),
                    "other": r(60),
                },
            },
        },
        "analysis_of_raw_data": {
            "base_calling": {
                "base_calling_software": {
                    "name_and_version": r(63),
                    "quality_score_version": r(64),
                    "other": r(65),
                },
                "output_files": {
                    "naming_convention": r(67),
                    "experimental_metadata_file": r(68),
                    "other": r(69),
                },
            },
            "availability_of_raw_data": {
                "output_files": {
                    "name_of_public_repository": r(72),
                    "accession_number": r(73),
                    "format_of_submitted_data": r(74),
                },
            },
            "raw_data_filtering": {
                "data_filtering_and_trimming": {
                    "minimum_read_count": r(77),
                    "list_of_samples_excluded_and_rationale": r(78),
                    "trimming_algorithm_software": r(79),
                    "trimming_parameters": r(80),
                    "other": r(81),
                },
            },
            "sequence_alignment": {
                "alignment": {
                    "source_and_version_of_genome_reference": r(84),
                    "mapping_software": r(85),
                    "mapping_parameters": r(86),
                    "other_relevant_information": r(87),
                },
            },
            "gene_quantification": {
                "quantification": {
                    "source_and_version_of_software": r(90),
                    "accession_number": r(91),
                    "format_of_submitted_data": r(92),
                    "method_for_summarizing_transcript_or_probe_counts": r(93),
                    "other": r(94),
                },
            },
        },
        "data_normalisation": {
            "normalisation_of_raw_count": {
                "normalisation_method_applied": r(96),
                "package_and_version_used": r(97),
                "factor_used_as_design": r(98),
                "other": r(99),
            },
        },
        "post_normalisation_data_filtering": {
            "outlier_and_low_quality_data_removal": {
                "outlier_identification_and_reason_for_removal": r(101),
                "thresholds": r(102),
                "processing_steps_where_exclusion_occurred": r(103),
                "list_of_samples_excluded_and_justification": r(104),
                "other": r(105),
            },
            "specific_gene_filtering": {
                "specific_gene_filtering": {
                    "low_read_count_filtering_step": r(108),
                    "high_variation_among_replicates_filtering_step": r(109),
                    "other_post_normalisation_gene_filtering_steps": r(110),
                    "other": r(111),
                },
            },
        },
    })


def parse_darm_dam(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Parse the DARM - DAM sheet."""
    r = lambda row: get_row_input(ws, row, 4)
    return prune_empty({
        "type": "DAMDARM",
        "inputs": {
            "data_input": r(6),
            "metadata_input": r(7),
        },
        "software_documentation": {
            "software": r(8),
            "operating_system": r(9),
            "additional_libraries_used": r(10),
            "software_availability": r(11),
        },
        "contrasts": {
            "contrasts": r(13),
        },
        "assay_experimental_design": {
            "group_sizes": r(15),
            "covariance": r(16),
            "technical_replicates": r(17),
        },
        "statistical_analysis": {
            "statistical_approach": r(19),
            "data_transformation": r(20),
            "effects_models": r(21),
            "modeling_inputs": r(22),
            "bayesian_approaches": r(23),
            "decision_criteria": r(24),
            "other": r(25),
        },
        "outputs": {
            "output_and_supporting_files": r(27),
        },
    })


def parse_darm_mva(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Parse the DARM - MVA sheet."""
    r = lambda row: get_row_input(ws, row, 4)
    return prune_empty({
        "type": "MVADARM",
        "inputs": {
            "data_input": r(6),
            "metadata_input": r(7),
        },
        "statistical_software": {
            "software": r(8),
            "operating_system": r(9),
            "additional_libraries_used": r(10),
            "software_availability": r(11),
        },
        "experimental_conditions": {
            "factors_and_levels_table": r(13),
        },
        "assay_experimental_design": {
            "group_sizes": r(15),
            "covariance": r(16),
            "technical_replicates": r(17),
        },
        "unsupervised_analysis": {
            "name_and_description_of_method": r(19),
            "groups_or_conditions_included": r(20),
            "variable_selection_method": r(21),
            "selected_variables": r(22),
            "parameter_settings": r(23),
            "other_or_citations": r(24),
        },
        "supervised_analysis": {
            "name_and_description_of_method": r(26),
            "groups_conditions_or_factors_contrasted": r(27),
            "variable_selection_method": r(28),
            "selected_variables": r(29),
            "parameter_settings": r(30),
            "citations_to_literature": r(31),
        },
        "validation": {
            "name_and_description_of_validation_method": r(33),
            "parameter_settings": r(34),
            "citations_to_literature": r(36),
        },
        "variable_importance_for_feature_selection": {
            "name_and_description_of_method": r(38),
            "details_of_how_significance_threshold_determined": r(39),
            "citations_to_literature": r(41),
        },
        "outputs": {
            "output_and_supporting_files": r(43),
        },
    })


def parse_darm_bmd(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, Any]:
    """Parse the DARM - BMD sheet."""
    r = lambda row: get_row_input(ws, row, 4)
    return prune_empty({
        "type": "BMDDARM",
        "inputs": {
            "data_input": r(6),
            "metadata_input": r(7),
        },
        "software_documentation": {
            "software": r(8),
            "operating_system": r(9),
            "additional_libraries_used": r(10),
            "software_availability": r(11),
        },
        "description_of_data_to_be_modelled": {
            "dose_levels": r(13),
            "annotation_of_model_substrate_data": r(14),
            "version_of_platform_annotation": r(15),
        },
        "pre_filtering_of_data": {
            "statistical_test_performed": r(17),
            "statistical_threshold_applied": r(18),
            "statistical_multiple_testing_correction_method": r(19),
            "additional_statistical_test_parameters": r(20),
            "additional_filtering": r(21),
        },
        "benchmark_dose_modeling": {
            "bmd_modeling_approach": r(23),
            "method_for_final_estimation_of_msd_bmd": r(24),
            "list_of_models_fit_to_data": r(25),
            "model_averaging_approach": r(26),
            "method_for_determining_bmr": r(27),
            "benchmark_response_bmr_factor": r(28),
            "number_of_fitting_iterations": r(29),
            "confidence_interval_of_bmd": r(30),
            "power_parameter_restriction": r(31),
            "variance_assumption": r(32),
            "model_selection": r(33),
            "criteria_to_identify_extrapolated_models": r(34),
        },
        "determination_of_biological_entity_or_set_level_bmd_values": {
            "biological_entity_or_set_annotation_used": r(36),
            "msd_acceptance_criteria": r(37),
            "criteria_for_identification_of_active_biological_sets": r(38),
            "method_of_estimating_set_bmd_bmdl_bmdu": r(39),
        },
        "outputs": {
            "output_and_supporting_files": r(41),
        },
    })


def parse_file_manifest(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[dict[str, Any]]:
    """Parse the File Manifest sheet (rows starting at row 4)."""
    entries: list[dict[str, Any]] = []
    for row_idx in range(4, ws.max_row + 1):
        desc = clean_val(ws.cell(row=row_idx, column=1).value)
        fname = clean_val(ws.cell(row=row_idx, column=2).value)
        md5 = clean_val(ws.cell(row=row_idx, column=3).value)
        if desc or fname or md5:
            entries.append(
                prune_empty({
                    "file_description": desc,
                    "file_name": fname,
                    "md5_checksum": md5,
                })
            )
    return entries


DAPRM_PARSERS = {
    "DAPRM - RNA-Seq": parse_daprm_rna_seq,
}

DARM_PARSERS = {
    "DARM - DAM": parse_darm_dam,
    "DARM - MVA": parse_darm_mva,
    "DARM - BMD": parse_darm_bmd,
}


def convert_workbook(xlsx_path: Path) -> dict[str, Any]:
    """Convert an OORF Excel workbook into a dictionary valid against the LinkML OORF schema.

    SSRM, TERM and File Manifest must be present; DAPRM and DARM sheets are parsed only if the workbook contains them.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    oorf_instance: dict[str, Any] = {
        "ssrm": parse_ssrm(wb["SSRM"]),
        "term": parse_term(wb["TERM"]),
        "daprm": [parse(wb[name]) for name, parse in DAPRM_PARSERS.items() if name in wb.sheetnames],
        "darm": [parse(wb[name]) for name, parse in DARM_PARSERS.items() if name in wb.sheetnames],
        "file_manifest": parse_file_manifest(wb["File Manifest"]),
    }
    return oorf_instance


def write_instance(instance: dict[str, Any], out_path: Path) -> None:
    """Write the instance as JSON or YAML, depending on the file extension."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.suffix in (".yaml", ".yml"):
        text = yaml.safe_dump(instance, sort_keys=False, allow_unicode=True, width=120)
    else:
        text = json.dumps(instance, indent=2, ensure_ascii=False) + "\n"
    out_path.write_text(text, encoding="utf-8")
    print(f"Wrote {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path, help="filled-in OORF Excel workbook")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        action="append",
        help="output file (.json, .yaml or .yml); can be repeated. Default: the workbook path with a .json suffix",
    )
    args = parser.parse_args()

    instance = convert_workbook(args.xlsx)
    for out_path in args.output or [args.xlsx.with_suffix(".json")]:
        write_instance(instance, out_path)


if __name__ == "__main__":
    main()
